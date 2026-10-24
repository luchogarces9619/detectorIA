"""Generadores deterministas de documentos para las pruebas de extracción."""

from __future__ import annotations

import io

import pymupdf
from docx import Document as DocxDocument
from PIL import Image, ImageDraw


def build_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("Impacto de la inteligencia artificial en la educación superior", level=0)
    document.add_heading("Resumen", level=1)
    document.add_paragraph(
        "Este estudio analiza el uso de herramientas de inteligencia artificial "
        "en estudiantes universitarios de Ecuador durante el año 2025."
    )
    document.add_heading("Introducción", level=1)
    document.add_paragraph(
        "La inteligencia artificial generativa ha transformado las prácticas "
        "académicas y plantea nuevos retos para la integridad educativa."
    )

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Variable"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "Participantes"
    table.cell(1, 1).text = "120"

    document.add_heading("Referencias", level=1)
    document.add_paragraph(
        "Pérez, J. (2023). Inteligencia artificial y educación. Revista Académica, 12(3), 45-67."
    )

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_text_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Introduccion", fontsize=18)
    page.insert_text(
        (72, 110),
        "Este documento describe los resultados de una investigacion academica.",
        fontsize=11,
    )
    page.insert_text(
        (72, 130),
        "Los datos fueron recolectados durante el ano 2025 en varias universidades.",
        fontsize=11,
    )
    data = document.tobytes()
    document.close()
    return data


def build_scanned_pdf() -> bytes:
    image = Image.new("RGB", (900, 260), "white")
    draw = ImageDraw.Draw(image)
    draw.text((30, 100), "Texto escaneado de prueba para el OCR academico", fill="black")

    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")

    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(50, 50, 550, 250), stream=image_buffer.getvalue())
    data = document.tobytes()
    document.close()
    return data
